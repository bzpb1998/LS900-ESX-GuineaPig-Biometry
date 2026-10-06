from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ls900_biometry.exporter import export_report  # noqa: E402
from ls900_biometry.webapp import create_app  # noqa: E402


class ApplicationSmokeTests(unittest.TestCase):
    def test_home_page(self) -> None:
        app = create_app()
        app.testing = True
        with app.test_client() as client:
            response = client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"LS900 ESX Guinea Pig", response.data)

    def test_compact_report_layout(self) -> None:
        common = {
            "animal_id": "2",
            "source_file": "ZA_2.esx",
            "source_key": "1:ZA_2.esx",
            "predicted_AL_mm": 7.01,
            "AL_quality": "high",
            "predicted_CCT_mm": 0.183,
            "CCT_quality": "high",
            "predicted_RT_mm": 0.127,
        }
        predictions = [
            {
                **common,
                "eye": "OD",
                "predicted_AD_mm": 0.94,
                "predicted_LT_mm": 3.50,
                "lens_quality": "usable",
                "lens_gate_reasons": "",
            },
            {
                **common,
                "eye": "OS",
                "predicted_AL_mm": 7.05,
                "predicted_AD_mm": float("nan"),
                "predicted_LT_mm": float("nan"),
                "lens_quality": "unmeasurable_or_review",
                "lens_gate_reasons": "unstable_lens_front",
            },
        ]
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "report.xlsx"
            export_report(predictions, target, title="LS900_Biometry_test")
            workbook = load_workbook(target)
            self.assertEqual(workbook.sheetnames, ["测量数据", "说明"])
            sheet = workbook["测量数据"]
            self.assertEqual(sheet["B6"].value, "2")
            self.assertEqual(sheet["C6"].value, 7.01)
            self.assertEqual(sheet["D6"].value, 183)
            self.assertEqual(sheet["E6"].value, 0.94)
            self.assertEqual(sheet["F6"].value, 3.50)
            self.assertEqual(sheet["H6"].value, None)
            self.assertEqual(sheet["I6"].value, None)
            self.assertEqual(sheet["J6"].value, 7.05)
            self.assertEqual(sheet["L6"].value, None)
            self.assertEqual(sheet["M6"].value, None)


if __name__ == "__main__":
    unittest.main()
