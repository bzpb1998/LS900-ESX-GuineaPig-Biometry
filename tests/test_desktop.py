"""Synthetic-only shared service and native UI regression tests."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook
from ls900_biometry.analysis import MODEL_VERSION, load_model
from ls900_biometry.batch import AnalysisCancelled, BatchResult, run_batch, save_report
from ls900_biometry.desktop import MainWindow


class DesktopTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_default_model_is_v5(self):
        self.assertIn("v5.0.0", MODEL_VERSION)
        self.assertEqual(load_model()["version"], MODEL_VERSION)

    def test_window_has_native_workflow(self):
        window = MainWindow()
        self.assertEqual(window.pages.count(), 3)
        self.assertFalse(window.save_button.isEnabled())
        window.navigate(2)
        self.assertEqual(window.pages.currentIndex(), 2)
        window.close()

    def test_duplicate_files_are_not_added_twice(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "synthetic.ESX"
            path.touch()
            window = MainWindow()
            window.add_paths([path, path, Path(temp)])
            self.assertEqual(len(window.paths), 1)
            self.assertTrue(window.analyze_button.isEnabled())
            window.clear_files()
            self.assertFalse(window.analyze_button.isEnabled())
            window.close()

    def test_bad_file_does_not_drop_other_files(self):
        with patch("ls900_biometry.batch.analyze_esx", side_effect=[ValueError("bad XML"), [{"predicted_AL_mm":7.2}]]):
            result = run_batch([Path("bad.esx"), Path("good.esx")])
        self.assertEqual(result.completed, 1)
        self.assertEqual(result.usable_eyes, 1)
        self.assertEqual(len(result.errors), 1)

    def test_cancellation_stops_batch(self):
        with self.assertRaises(AnalysisCancelled):
            run_batch([Path("synthetic.esx")], cancelled=lambda: True)

    def test_export_has_v5_metadata_and_keeps_template(self):
        row = dict(animal_id="1", eye="OD", source_file="synthetic.esx", source_key="synthetic",
                   algorithm_version=MODEL_VERSION, predicted_AL_mm=7.2, predicted_CCT_mm=.18,
                   predicted_AD_mm=.95, predicted_LT_mm=3.5, predicted_RT_mm=.13,
                   AL_quality="high", CCT_quality="high", lens_quality="unmeasurable_or_review")
        with tempfile.TemporaryDirectory() as temp:
            output = save_report(BatchResult(rows=[row], completed=1, total=1), Path(temp)/"test.xlsx")
            book = load_workbook(output, data_only=True)
            self.assertEqual(book.sheetnames, ["测量数据","说明"])
            sheet = book.worksheets[0]
            self.assertEqual(sheet["C6"].value,7.2)
            self.assertEqual(sheet["D6"].value,180)
            self.assertTrue(all(sheet[x].value is None for x in ("E6","F6","H6","I6","J6")))
            notes = " ".join(str(book.worksheets[1].cell(r,2).value) for r in range(4,book.worksheets[1].max_row+1))
            self.assertIn(MODEL_VERSION, notes)
            self.assertNotIn("锁定的 V4", notes)

    def test_failed_batch_cannot_produce_empty_report(self):
        with self.assertRaises(ValueError):
            save_report(BatchResult(), Path("must_not_exist.xlsx"))


if __name__ == "__main__":
    unittest.main()
