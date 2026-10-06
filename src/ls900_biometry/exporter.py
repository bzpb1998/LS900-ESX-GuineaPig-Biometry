"""Create the compact two-sheet Excel report requested for routine use."""

from __future__ import annotations

import math
from collections import Counter, defaultdict
from copy import copy
from importlib.resources import files
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


TEMPLATE_PATH = files("ls900_biometry").joinpath("assets/output_template.xlsx")
REVIEW_FILL = PatternFill("solid", fgColor="FFC7CE")
BLANK_FILL = PatternFill("solid", fgColor="F2F2F2")
REASON_LABELS = {
    "no_usable_ascan": "没有可用的晶体 A-scan",
    "low_paired_peak_fraction": "晶体前后表面成对峰比例偏低",
    "unstable_lens_front": "晶体前表面峰不稳定",
    "unstable_lens_back": "晶体后表面峰不稳定",
    "dual_algorithm_disagreement": "两种晶体算法结果分歧较大",
    "LT_out_of_range_or_missing": "晶体厚度缺失或超出训练范围",
    "eye_processing_error": "该眼处理失败",
}


def _finite(value: Any) -> bool:
    try:
        return math.isfinite(float(value))
    except (TypeError, ValueError):
        return False


def _copy_row_style(sheet, source_row: int, target_row: int) -> None:
    for column in range(1, 15):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def _natural_key(label: str) -> tuple[int, int | str]:
    return (0, int(label)) if label.isdigit() else (1, label.lower())


def export_report(
    predictions: Iterable[dict[str, Any]],
    output_path: str | Path,
    *,
    title: str,
    extra_notes: Iterable[str] = (),
    file_errors: Iterable[str] = (),
) -> Path:
    """Write one animal per row, OD on the left and OS on the right."""

    rows = list(predictions)
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    file_sequence: list[tuple[str, str, str]] = []
    seen_files: set[tuple[str, str]] = set()
    for item in rows:
        source_file = str(item.get("source_file", ""))
        source_key = str(item.get("source_key", source_file))
        file_key = (source_key, str(item.get("animal_id", "")))
        if file_key not in seen_files:
            seen_files.add(file_key)
            file_sequence.append((source_key, file_key[1], source_file))
        grouped[(source_key, str(item.get("eye", "")))] = item

    file_sequence.sort(key=lambda item: (_natural_key(item[1]), item[0].lower()))
    id_counts = Counter(animal_id for _, animal_id, _ in file_sequence)
    id_seen: defaultdict[str, int] = defaultdict(int)

    workbook = load_workbook(TEMPLATE_PATH)
    data = workbook["测量数据"]
    notes = workbook["说明"]
    data["B2"] = title

    required_rows = max(1, len(file_sequence))
    for offset in range(1, required_rows):
        _copy_row_style(data, 6, 6 + offset)
    if data.max_row > 5 + required_rows:
        data.delete_rows(6 + required_rows, data.max_row - (5 + required_rows))

    explanation: list[str] = [
        "本表由锁定的 V4 模型直接解析原始 ESX 文件生成。",
        "AD、LT 仅在晶体质量门通过时显示；未通过时留空，并在本页列明原因。",
        "RT 为探索性结果，尚未作为主要终点完成独立验证，使用时请谨慎。",
    ]

    metric_map = [
        ("predicted_AL_mm", 3, "0.00"),
        ("predicted_CCT_mm", 4, "0"),
        ("predicted_AD_mm", 5, "0.00"),
        ("predicted_LT_mm", 6, "0.00"),
        ("predicted_RT_mm", 7, "0"),
    ]
    for index, (source_key, animal_id, source_file) in enumerate(file_sequence, start=6):
        id_seen[animal_id] += 1
        display_id = animal_id
        if id_counts[animal_id] > 1:
            display_id = f"{animal_id}-{id_seen[animal_id]}"
            explanation.append(f"编号 {display_id} 对应文件：{source_file}")
        data.cell(index, 2, display_id)

        for eye, base_column in (("OD", 3), ("OS", 10)):
            item = grouped.get((source_key, eye))
            if item is None:
                for column in range(base_column, base_column + 5):
                    data.cell(index, column).fill = BLANK_FILL
                explanation.append(f"{display_id} {eye}：文件中没有可输出的数据。")
                continue

            for metric, relative_column, number_format in metric_map:
                column = base_column + relative_column - 3
                cell = data.cell(index, column)
                value = item.get(metric)
                cell.number_format = number_format
                if _finite(value):
                    numeric = float(value)
                    cell.value = numeric * 1000.0 if metric in {"predicted_CCT_mm", "predicted_RT_mm"} else numeric
                else:
                    cell.value = None
                    cell.fill = BLANK_FILL

            if item.get("AL_quality") == "review":
                data.cell(index, base_column).fill = REVIEW_FILL
                explanation.append(
                    f"{display_id} {eye}：眼轴峰位选择存在较大分歧，数值已保留但建议结合原始波形复核。"
                )
            if item.get("CCT_quality") == "review":
                data.cell(index, base_column + 1).fill = REVIEW_FILL
                explanation.append(f"{display_id} {eye}：角膜厚度超出当前自动质量范围。")
            if item.get("lens_quality") != "usable":
                data.cell(index, base_column + 2).fill = BLANK_FILL
                data.cell(index, base_column + 3).fill = BLANK_FILL
                raw_reasons = str(item.get("lens_gate_reasons", ""))
                translated = [REASON_LABELS.get(reason, reason) for reason in raw_reasons.split(";") if reason]
                reason_text = "；".join(translated) or "晶体结果未通过质量门"
                explanation.append(f"{display_id} {eye}：AD/LT 留空（{reason_text}）。")
            if item.get("processing_error"):
                explanation.append(f"{display_id} {eye}：处理错误：{item['processing_error']}")

    explanation.extend(str(note) for note in extra_notes if str(note).strip())
    explanation.extend(f"文件处理失败：{error}" for error in file_errors)
    notes["B2"] = "说明"
    for merged_range in list(notes.merged_cells.ranges):
        if str(merged_range) != "B2:N2":
            notes.unmerge_cells(str(merged_range))
    if notes.max_row > 3:
        notes.delete_rows(4, notes.max_row - 3)
    for row_index, line in enumerate(explanation, start=4):
        notes.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=14)
        notes.cell(row_index, 2, f"• {line}")
        notes.cell(row_index, 2).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        notes.cell(row_index, 2).font = Font(name="Microsoft YaHei", size=11, color="1B2A4A")
        notes.cell(row_index, 2).fill = PatternFill(
            "solid", fgColor="EAF0F8" if row_index % 2 == 0 else "FFFFFF"
        )
        notes.row_dimensions[row_index].height = 22

    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target
